"""Отчёт по использованию лимита консультаций: звонки и задачи ЭПД."""

import base64
import re
import tempfile
from datetime import date, datetime, timedelta
from pathlib import Path

import dateutil.parser
import openpyxl
from fast_bitrix24 import Bitrix
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from web_app_4dk.modules.authentication import authentication

if __package__:
    from .worklog_common import API, accounting_timezone, config
    from .sync_worklog import (
        COMPANY_EXCLUDE_FIELD,
        EPD_GROUP_ID,
        EPD_TAG,
        EXCLUDED_STAGE_IDS,
    )
else:
    from worklog_common import API, accounting_timezone, config
    from sync_worklog import (
        COMPANY_EXCLUDE_FIELD,
        EPD_GROUP_ID,
        EPD_TAG,
        EXCLUDED_STAGE_IDS,
    )


webhook = authentication('Bitrix')
b = Bitrix(webhook)

MONTH_CODES = {
    'Январь': 1,
    'Февраль': 2,
    'Март': 3,
    'Апрель': 4,
    'Май': 5,
    'Июнь': 6,
    'Июль': 7,
    'Август': 8,
    'Сентябрь': 9,
    'Октябрь': 10,
    'Ноябрь': 11,
    'Декабрь': 12,
}

CONSULTATION_DEPARTMENT_ID = '231'
EXCLUDED_CALL_AUTHOR_IDS = {'109', '19'}
EXCLUDED_WORKLOG_EMPLOYEE_ID = '173'
BITRIX_REPORT_FOLDER_ID = '187139'
TIMELINE_AUTHOR_ID = '173'

THIN_BORDER = Border(
    left=Side(style='thin', color='D9E2F3'),
    right=Side(style='thin', color='D9E2F3'),
    top=Side(style='thin', color='D9E2F3'),
    bottom=Side(style='thin', color='D9E2F3'),
)

HEADER_FILL = PatternFill(fill_type='solid', fgColor='2F5597')
SECTION_FILL = PatternFill(fill_type='solid', fgColor='D9EAF7')
TOTAL_FILL = PatternFill(fill_type='solid', fgColor='E2F0D9')


def _request_period(req):
    month_name = str(req['month'])
    year = str(req['year'])

    if month_name not in MONTH_CODES:
        raise ValueError('Неизвестный месяц: ' + month_name)

    if not re.fullmatch(r'\d{4}', year):
        raise ValueError('Некорректный год: ' + year)

    month_number = MONTH_CODES[month_name]
    period = f'{year}-{month_number:02d}'

    start = datetime(int(year), month_number, 1)
    if month_number == 12:
        end = datetime(int(year) + 1, 1, 1)
    else:
        end = datetime(int(year), month_number + 1, 1)

    return month_name, year, period, start, end


def _get_company(company_id):
    company = b.get_all('crm.company.get', {'ID': company_id})

    if isinstance(company, list):
        company = company[0] if company else None

    if not isinstance(company, dict):
        raise RuntimeError(f'Не удалось получить компанию {company_id}')

    return company


def _get_user(user_id, cache):
    user_id = str(user_id)

    if user_id not in cache:
        rows = b.get_all('user.get', {'ID': user_id})
        cache[user_id] = rows[0] if rows else None

    return cache[user_id]


def _user_name(user_id, cache):
    user = _get_user(user_id, cache)

    if not user:
        return f'Пользователь {user_id}'

    parts = [
        str(user.get('NAME') or '').strip(),
        str(user.get('LAST_NAME') or '').strip(),
    ]
    return ' '.join(part for part in parts if part) or f'Пользователь {user_id}'


def _contact_name(contact_id):
    contact = b.get_all('crm.contact.get', {'ID': contact_id})

    if isinstance(contact, list):
        contact = contact[0] if contact else None

    if not isinstance(contact, dict):
        return f'Контакт {contact_id}'

    parts = [
        str(contact.get('LAST_NAME') or '').strip(),
        str(contact.get('NAME') or '').strip(),
        str(contact.get('SECOND_NAME') or '').strip(),
    ]
    return ' '.join(part for part in parts if part)


def _call_phone(subject):
    subject = str(subject or '')

    if ' на ' not in subject:
        return ''

    return subject.split(' на ', 1)[1].strip()


def _load_owner_calls(
    owner_type_id,
    owner_id,
    contact_name,
    period,
    range_start,
    range_end,
    user_cache,
    seen_activity_ids,
):
    rows = []

    activities = b.get_all(
        'crm.activity.list',
        {
            'select': [
                'ID',
                'SUBJECT',
                'START_TIME',
                'END_TIME',
                'AUTHOR_ID',
            ],
            'filter': {
                'OWNER_TYPE_ID': str(owner_type_id),
                'OWNER_ID': str(owner_id),
                'PROVIDER_TYPE_ID': 'CALL',
                '>=START_TIME': range_start.isoformat(),
                '<START_TIME': range_end.isoformat(),
            },
        },
    )

    for activity in activities:
        activity_id = str(activity.get('ID') or '')

        if not activity_id or activity_id in seen_activity_ids:
            continue

        subject = str(activity.get('SUBJECT') or '')
        if 'Исходящий' not in subject:
            continue

        author_id = str(activity.get('AUTHOR_ID') or '')
        if author_id in EXCLUDED_CALL_AUTHOR_IDS:
            continue

        author = _get_user(author_id, user_cache)
        if not author:
            raise RuntimeError(
                f'Звонок {activity_id}: не найден пользователь {author_id}'
            )

        departments = {
            str(value)
            for value in (author.get('UF_DEPARTMENT') or [])
        }
        if CONSULTATION_DEPARTMENT_ID not in departments:
            continue

        start_value = activity.get('START_TIME')
        end_value = activity.get('END_TIME')

        if not start_value or not end_value:
            continue

        call_start = dateutil.parser.isoparse(start_value)
        call_end = dateutil.parser.isoparse(end_value)

        # Дополнительная защита на случай особенностей фильтра Битрикс24.
        if call_start.strftime('%Y-%m') != period:
            continue

        duration = call_end - call_start
        if duration.total_seconds() < 0:
            raise ValueError(
                f'Звонок {activity_id}: время окончания раньше времени начала'
            )

        seen_activity_ids.add(activity_id)

        rows.append(
            {
                'date': call_start,
                'date_text': call_start.strftime('%d.%m.%Y %H:%M:%S'),
                'contact': contact_name,
                'phone': _call_phone(subject),
                'duration': duration,
                'employee': _user_name(author_id, user_cache),
            }
        )

    return rows


def _collect_calls(
    company_id,
    contacts,
    deals,
    period,
    range_start,
    range_end,
    user_cache,
):
    result = []
    seen_activity_ids = set()

    for contact in contacts:
        contact_id = contact.get('CONTACT_ID')
        if not contact_id:
            continue

        result.extend(
            _load_owner_calls(
                owner_type_id=3,
                owner_id=contact_id,
                contact_name=_contact_name(contact_id),
                period=period,
                range_start=range_start,
                range_end=range_end,
                user_cache=user_cache,
                seen_activity_ids=seen_activity_ids,
            )
        )

    for deal in deals:
        deal_id = deal.get('ID')
        if not deal_id:
            continue

        result.extend(
            _load_owner_calls(
                owner_type_id=2,
                owner_id=deal_id,
                contact_name='',
                period=period,
                range_start=range_start,
                range_end=range_end,
                user_cache=user_cache,
                seen_activity_ids=seen_activity_ids,
            )
        )

    return sorted(result, key=lambda row: row['date'])


def _field_matches(value, expected):
    return str(value or '') == str(expected)


def _journal_datetime(value, settings):
    if not value:
        raise ValueError('В записи журнала отсутствует дата трудозатраты')

    dt = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
    tz = accounting_timezone(settings['timezone'])

    if dt.tzinfo is None:
        return dt.replace(tzinfo=tz)

    return dt.astimezone(tz)


def _company_excludes_epd(company):
    if COMPANY_EXCLUDE_FIELD not in company:
        raise RuntimeError(
            f'В компании не получено поле {COMPANY_EXCLUDE_FIELD}'
        )

    value = company[COMPANY_EXCLUDE_FIELD]

    if isinstance(value, list):
        if len(value) > 1:
            raise RuntimeError(
                f'Некорректное значение поля {COMPANY_EXCLUDE_FIELD}: {value}'
            )
        value = value[0] if value else None

    if value is None or value is False or value == '':
        return False

    if type(value) is int:
        if value == 0:
            return False
        if value == 1:
            return True

    if isinstance(value, str):
        normalized = value.strip().upper()

        if normalized in {'0', 'N', 'FALSE'}:
            return False

        if normalized in {'1', 'Y', 'TRUE'}:
            return True

    if value is True:
        return True

    raise RuntimeError(
        f'Некорректное значение поля {COMPANY_EXCLUDE_FIELD}: {value!r}'
    )


def _load_tasks(api, task_ids):
    result = {}

    for offset in range(0, len(task_ids), 50):
        chunk = task_ids[offset:offset + 50]

        tasks = api.pages(
            'tasks.task.list',
            {
                'filter': {'ID': [int(task_id) for task_id in chunk]},
                'order': {'ID': 'ASC'},
                'select': ['ID', 'TITLE', 'STAGE_ID'],
            },
            'tasks',
        )

        for task in tasks:
            task_id = str(task.get('id') or task.get('ID') or '')
            if task_id:
                result[task_id] = task

    missing = sorted(set(task_ids) - set(result), key=int)
    if missing:
        raise RuntimeError(
            'Не удалось получить задачи: ' + ', '.join(missing)
        )

    return result


def _task_stage(task):
    value = task.get('stageId')
    if value is None:
        value = task.get('STAGE_ID')

    if value is None or not re.fullmatch(r'\d+', str(value)):
        raise RuntimeError(
            f'Задача {task.get("id") or task.get("ID")}: '
            f'не получена корректная стадия'
        )

    return int(value)


def _task_title(task, task_id):
    return str(
        task.get('title')
        or task.get('TITLE')
        or f'Задача {task_id}'
    )


def _collect_epd_worklogs(
    company_id,
    company,
    period,
    user_cache,
):
    settings = config()
    api = API(settings)
    api.validate_fields()

    fields = settings['fields']

    journal_rows = api.pages(
        'crm.item.list',
        {
            'entityTypeId': settings['entity_type_id'],
            'useOriginalUfNames': 'Y',
            'filter': {
                'categoryId': settings['category_id'],
                'companyId': int(company_id),
                fields['period']: period,
            },
            'order': {'id': 'ASC'},
            'select': ['*'],
        },
        'items',
    )

    candidates = []
    seen_keys = set()
    accounting_start = date.fromisoformat(settings['start_date'])

    for row in journal_rows:
        if str(row.get('companyId') or '') != str(company_id):
            continue

        if not _field_matches(
            row.get(fields['source']),
            settings['source_id'],
        ):
            continue

        if not _field_matches(
            row.get(fields['group']),
            EPD_GROUP_ID,
        ):
            continue

        if str(row.get(fields['tag']) or '') != EPD_TAG:
            continue

        if not _field_matches(
            row.get(fields['kind']),
            settings['general_kind'],
        ):
            continue

        if not _field_matches(
            row.get(fields['state']),
            settings['active_state'],
        ):
            continue

        employee_id = str(row.get(fields['employee']) or '')
        if employee_id == EXCLUDED_WORKLOG_EMPLOYEE_ID:
            continue

        key = str(row.get(fields['key']) or '')
        if not key:
            raise RuntimeError(
                f'Элемент журнала {row.get("id")} не содержит ключ'
            )

        if key in seen_keys:
            raise RuntimeError(f'Дубли записей журнала по ключу {key}')
        seen_keys.add(key)

        task_id = str(row.get(fields['task']) or '')
        if not re.fullmatch(r'\d+', task_id):
            raise RuntimeError(
                f'Элемент журнала {row.get("id")}: '
                f'некорректный ID задачи'
            )

        work_date = _journal_datetime(
            row.get(fields['date']),
            settings,
        )

        if work_date.strftime('%Y-%m') != period:
            continue

        if work_date.date() < accounting_start:
            continue

        seconds = int(row.get(fields['seconds']) or 0)
        if seconds < 0:
            raise RuntimeError(
                f'Элемент журнала {row.get("id")}: '
                f'отрицательная продолжительность'
            )

        if seconds == 0:
            continue

        candidates.append(
            {
                'journal_id': row.get('id'),
                'task_id': task_id,
                'entry_id': row.get(fields['entry']),
                'date': work_date,
                'description': str(
                    row.get(fields['description']) or ''
                ).strip(),
                'seconds': seconds,
                'employee_id': employee_id,
            }
        )

    if not candidates:
        return []

    # Если компания целиком исключена из списания ЭПД,
    # её трудозатраты не должны попадать в расход лимита.
    if _company_excludes_epd(company):
        return []

    task_ids = sorted(
        {row['task_id'] for row in candidates},
        key=int,
    )
    tasks = _load_tasks(api, task_ids)

    result = []

    for row in candidates:
        task = tasks[row['task_id']]

        # Такое же исключение, как в aggregate() сборщика.
        if _task_stage(task) in EXCLUDED_STAGE_IDS:
            continue

        result.append(
            {
                'date': row['date'],
                'date_text': row['date'].strftime('%d.%m.%Y %H:%M:%S'),
                'task_id': row['task_id'],
                'task_title': _task_title(task, row['task_id']),
                'description': row['description'],
                'duration': timedelta(seconds=row['seconds']),
                'employee': _user_name(
                    row['employee_id'],
                    user_cache,
                ),
            }
        )

    return sorted(
        result,
        key=lambda row: (
            row['date'],
            int(row['task_id']),
        ),
    )


def _safe_filename_part(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value))
    value = re.sub(r'\s+', '_', value.strip())
    return value[:100] or 'Компания'


def _style_row(worklist, row_number, fill, white_font=False):
    for cell in worklist[row_number]:
        cell.fill = fill
        cell.font = Font(
            bold=True,
            color='FFFFFF' if white_font else '000000',
        )
        cell.alignment = Alignment(
            vertical='center',
            wrap_text=True,
        )


def _format_workbook(worklist, header_rows, section_rows, total_rows):
    for row_number in header_rows:
        _style_row(
            worklist,
            row_number,
            HEADER_FILL,
            white_font=True,
        )

    for row_number in section_rows:
        _style_row(
            worklist,
            row_number,
            SECTION_FILL,
        )

    for row_number in total_rows:
        _style_row(
            worklist,
            row_number,
            TOTAL_FILL,
        )

    for row in worklist.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical='top',
                wrap_text=True,
            )

            if isinstance(cell.value, timedelta):
                cell.number_format = '[h]:mm:ss'

            if cell.value is not None:
                cell.border = THIN_BORDER

    worklist.freeze_panes = 'A4'
    worklist.sheet_view.showGridLines = False
    worklist.page_setup.orientation = 'landscape'
    worklist.page_setup.fitToWidth = 1
    worklist.page_setup.fitToHeight = 0
    worklist.sheet_properties.pageSetUpPr.fitToPage = True

    preferred_widths = {
        1: 22,
        2: 25,
        3: 42,
        4: 42,
        5: 18,
        6: 26,
    }

    for column_number, width in preferred_widths.items():
        worklist.column_dimensions[
            get_column_letter(column_number)
        ].width = width


def _create_workbook(
    company_name,
    month_name,
    year,
    report_created_time,
    calls,
    task_rows,
):
    workbook = openpyxl.Workbook()
    worklist = workbook.active
    worklist.title = 'Расход лимита'

    call_total = sum(
        (row['duration'] for row in calls),
        timedelta(),
    )
    task_total = sum(
        (row['duration'] for row in task_rows),
        timedelta(),
    )
    overall_total = call_total + task_total

    header_rows = []
    section_rows = []
    total_rows = []

    worklist['A1'] = company_name
    worklist['C1'] = f'{month_name} {year}'
    worklist['E1'] = (
        'Дата формирования '
        + report_created_time.strftime('%d.%m.%Y %H:%M')
    )

    worklist.merge_cells('A1:B1')
    worklist.merge_cells('C1:D1')
    worklist.merge_cells('E1:F1')

    _style_row(worklist, 1, HEADER_FILL, white_font=True)
    worklist.row_dimensions[1].height = 34

    worklist.append([])

    summary_title_row = worklist.max_row + 1
    worklist.append(['Использование лимита за период'])
    worklist.merge_cells(
        start_row=summary_title_row,
        start_column=1,
        end_row=summary_title_row,
        end_column=6,
    )
    section_rows.append(summary_title_row)

    summary_header_row = worklist.max_row + 1
    worklist.append(['Вид расхода', 'Продолжительность'])
    header_rows.append(summary_header_row)

    worklist.append(['Исходящие звонки', call_total])
    worklist.append(['Работы по задачам ЭПД', task_total])

    summary_total_row = worklist.max_row + 1
    worklist.append(['Всего израсходовано', overall_total])
    total_rows.append(summary_total_row)

    worklist.append([])

    calls_title_row = worklist.max_row + 1
    worklist.append(['Исходящие звонки линии консультаций'])
    worklist.merge_cells(
        start_row=calls_title_row,
        start_column=1,
        end_row=calls_title_row,
        end_column=6,
    )
    section_rows.append(calls_title_row)

    calls_header_row = worklist.max_row + 1
    worklist.append(
        [
            'Дата звонка',
            'Контакт (кому звонили)',
            'Номер телефона, на который звонили',
            'Хронометраж',
            'Кто звонил',
        ]
    )
    header_rows.append(calls_header_row)

    if calls:
        for row in calls:
            worklist.append(
                [
                    row['date_text'],
                    row['contact'],
                    row['phone'],
                    row['duration'],
                    row['employee'],
                ]
            )
    else:
        empty_row = worklist.max_row + 1
        worklist.append(
            ['За выбранный период учтённых звонков нет']
        )
        worklist.merge_cells(
            start_row=empty_row,
            start_column=1,
            end_row=empty_row,
            end_column=5,
        )

    calls_total_row = worklist.max_row + 1
    worklist.append(['Итого по звонкам', '', '', call_total])
    total_rows.append(calls_total_row)

    worklist.append([])

    tasks_title_row = worklist.max_row + 1
    worklist.append(['Работы по задачам ЭПД'])
    worklist.merge_cells(
        start_row=tasks_title_row,
        start_column=1,
        end_row=tasks_title_row,
        end_column=6,
    )
    section_rows.append(tasks_title_row)

    tasks_header_row = worklist.max_row + 1
    worklist.append(
        [
            'Дата работы',
            '№ задачи',
            'Название задачи',
            'Комментарий к трудозатрате',
            'Хронометраж',
            'Специалист',
        ]
    )
    header_rows.append(tasks_header_row)

    if task_rows:
        for row in task_rows:
            worklist.append(
                [
                    row['date_text'],
                    row['task_id'],
                    row['task_title'],
                    row['description'],
                    row['duration'],
                    row['employee'],
                ]
            )
    else:
        empty_row = worklist.max_row + 1
        worklist.append(
            [
                'За выбранный период учтённых '
                'трудозатрат по задачам ЭПД нет'
            ]
        )
        worklist.merge_cells(
            start_row=empty_row,
            start_column=1,
            end_row=empty_row,
            end_column=6,
        )

    tasks_total_row = worklist.max_row + 1
    worklist.append(
        ['Итого по задачам ЭПД', '', '', '', task_total]
    )
    total_rows.append(tasks_total_row)

    _format_workbook(
        worklist,
        header_rows=header_rows,
        section_rows=section_rows,
        total_rows=total_rows,
    )

    return workbook


def create_company_call_report(req):
    report_created_time = datetime.now()
    company_id = str(req['id'])

    month_name, year, period, range_start, range_end = (
        _request_period(req)
    )

    company = _get_company(company_id)
    company_name = str(company.get('TITLE') or f'Компания {company_id}')

    contacts = b.get_all(
        'crm.company.contact.items.get',
        {'id': company_id},
    )

    deals = b.get_all(
        'crm.deal.list',
        {
            'select': ['ID'],
            'filter': {'COMPANY_ID': company_id},
        },
    )

    user_cache = {}

    calls = _collect_calls(
        company_id=company_id,
        contacts=contacts,
        deals=deals,
        period=period,
        range_start=range_start,
        range_end=range_end,
        user_cache=user_cache,
    )

    task_rows = _collect_epd_worklogs(
        company_id=company_id,
        company=company,
        period=period,
        user_cache=user_cache,
    )

    workbook = _create_workbook(
        company_name=company_name,
        month_name=month_name,
        year=year,
        report_created_time=report_created_time,
        calls=calls,
        task_rows=task_rows,
    )

    timestamp = report_created_time.strftime(
        '%d-%m-%Y_%H-%M-%S_%f'
    )
    safe_company_name = _safe_filename_part(company_name)
    report_name = (
        f'Отчет_по_использованию_лимита_'
        f'{safe_company_name}_{timestamp}.xlsx'
    )

    with tempfile.TemporaryDirectory(
        prefix='company_limit_report_'
    ) as temp_directory:
        report_path = Path(temp_directory) / report_name
        workbook.save(report_path)
        workbook.close()

        with report_path.open('rb') as file:
            report_file_base64 = base64.b64encode(
                file.read()
            ).decode('ascii')

        upload_report = b.call(
            'disk.folder.uploadfile',
            {
                'id': BITRIX_REPORT_FOLDER_ID,
                'data': {'NAME': report_name},
                'fileContent': report_file_base64,
            },
        )

    if (
        isinstance(upload_report, dict)
        and isinstance(upload_report.get('result'), dict)
    ):
        upload_report = upload_report['result']

    detail_url = (
        upload_report.get('DETAIL_URL')
        or upload_report.get('detailUrl')
        if isinstance(upload_report, dict)
        else None
    )

    if not detail_url:
        raise RuntimeError(
            'Битрикс24 не вернул ссылку на загруженный отчёт'
        )

    b.call(
        'crm.timeline.comment.add',
        {
            'fields': {
                'ENTITY_ID': company_id,
                'ENTITY_TYPE': 'company',
                'COMMENT': (
                    f'Отчёт по использованию лимита консультаций '
                    f'за {month_name} {year}:\n'
                    f'{detail_url}'
                ),
                'AUTHOR_ID': TIMELINE_AUTHOR_ID,
            }
        },
    )

if __name__ == '__main__':
    req = {
        'id': '1973',
        'month': 'Сентябрь',
        'year': '2026',
    }

    create_company_call_report(req)