"""Отчёт по звонкам ЛК по контактам, выбранным при запуске БП.

Для БП вызывайте create_contact_consultation_calls_report(req), где req:
    {'month': 'Июль', 'year': '2026', 'contacts': ['C_5777', 'C_11579'],
     'user_id': 'user_1'}
Контакты можно передать и строкой '5777,+11579' через URL. user_id
необязателен: если его нет, файл загружается на Диск без уведомления.
Локальный запуск: python CreateContactConsultationCallsReport.py
Требуются fast-bitrix24, openpyxl и authentication.py
с функцией authentication('Bitrix') (либо пакет web_app_4dk.modules).

Критерий ЛК перенесён из CreateCompanyCallEpdReport.py: автор звонка
состоит в отделе 231, кроме пользователей 109, 19 и 117. Это не проверка
фактического номера, с которого был сделан вызов.
"""

import base64
import json
import re
import tempfile
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
from fast_bitrix24 import Bitrix
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

try:
    from authentication import authentication
except ImportError:
    from web_app_4dk.modules.authentication import authentication


CONSULTATION_DEPARTMENT_ID = '231'
EXCLUDED_CALL_AUTHOR_IDS = {'109', '19', '117'}
BITRIX_REPORT_FOLDER_ID = '213023'  # Та же папка, что в CreateLineConsultationReport.py
MONTH_NAMES = (
    'Январь', 'Февраль', 'Март', 'Апрель', 'Май', 'Июнь',
    'Июль', 'Август', 'Сентябрь', 'Октябрь', 'Ноябрь', 'Декабрь',
)


def parse_contact_ids(value):
    parts = [part.strip() for part in value.split(',')]
    if not parts or any(not re.fullmatch(r'[1-9]\d*', part) for part in parts):
        raise ValueError('Введите положительные ID контактов через запятую, например: 123, 456')
    return list(dict.fromkeys(parts))


def parse_bp_contact_ids(value):
    """Принимает множественную CRM-привязку БП или строку с ID через запятую."""
    if isinstance(value, str):
        value = value.strip()
        if value.startswith('['):
            try:
                value = json.loads(value)
            except json.JSONDecodeError as exc:
                raise ValueError('Контакты: неверная JSON-строка') from exc
        else:
            value = [part.strip() for part in value.split(',')]
    if not isinstance(value, (list, tuple)) or not value:
        raise ValueError('Выберите хотя бы один контакт при запуске БП')
    ids = []
    for item in value:
        if not isinstance(item, (str, int)):
            raise ValueError(f'Некорректное значение контакта: {item!r}')
        # В query string знак + обычно декодируется как пробел. Если маршрут
        # передаст его буквально, тоже принимаем перед положительным ID.
        match = re.fullmatch(r'(?:C_|\+)?([1-9]\d*)', str(item).strip())
        if not match:
            raise ValueError(f'Ожидался ID контакта или C_ID, получено: {item!r}')
        ids.append(match.group(1))
    return list(dict.fromkeys(ids))


def parse_period(value):
    match = re.fullmatch(r'\s*(0[1-9]|1[0-2])\s+(\d{4})\s*', value)
    if not match:
        raise ValueError('Введите месяц и год в формате 09 2026')
    month, year = map(int, match.groups())
    start = datetime(year, month, 1)
    end = datetime(year + (month == 12), month % 12 + 1, 1)
    return month, year, start, end


def parse_bp_period(req):
    month_value = str(req['month']).strip()
    year = str(req['year']).strip()
    if month_value in MONTH_NAMES:
        month_value = f'{MONTH_NAMES.index(month_value) + 1:02d}'
    elif month_value.isdigit():
        month_value = month_value.zfill(2)
    return parse_period(f'{month_value} {year}')


def parse_user_id(value):
    match = re.fullmatch(r'(?:user_)?([1-9]\d*)', str(value).strip())
    if not match:
        raise ValueError('Некорректный user_id запускающего БП')
    return int(match.group(1))


def get_record(bitrix, method, object_id):
    result = bitrix.get_all(method, {'ID': object_id})
    if isinstance(result, list):
        result = result[0] if result else None
    return result if isinstance(result, dict) else None


def contact_name(contact):
    return ' '.join(filter(None, (str(contact.get(field) or '').strip() for field in
                                  ('LAST_NAME', 'NAME', 'SECOND_NAME')))) or 'Не указан'


def user_name(user, user_id):
    return ' '.join(filter(None, (str(user.get(field) or '').strip() for field in
                                  ('NAME', 'LAST_NAME')))) or f'Пользователь {user_id}'


def call_phone(subject):
    subject = str(subject or '')
    return subject.split(' на ', 1)[1].strip() if ' на ' in subject else ''


def parse_bitrix_datetime(value):
    """Разбирает ISO дату Битрикс24 без стороннего пакета dateutil."""
    return datetime.fromisoformat(str(value).replace('Z', '+00:00'))


def collect_calls(bitrix, contact_ids, month, year, start, end):
    users = {}
    seen_ids = set()
    calls = []
    for contact_id in contact_ids:
        contact = get_record(bitrix, 'crm.contact.get', contact_id)
        if contact is None:
            raise RuntimeError(f'Не удалось получить контакт {contact_id}')
        name = contact_name(contact)
        # OWNER_ID отражает только основного владельца дела. Звонок может
        # отображаться в карточке контакта как дополнительная CRM-привязка,
        # даже если его основной владелец — сделка или компания.
        common_filter = {
            'PROVIDER_TYPE_ID': 'CALL',
            '>=START_TIME': start.isoformat(),
            '<START_TIME': end.isoformat(),
        }
        select = ['ID', 'SUBJECT', 'START_TIME', 'END_TIME', 'AUTHOR_ID']
        direct = bitrix.get_all('crm.activity.list', {
            'select': select,
            'filter': {**common_filter, 'OWNER_TYPE_ID': '3', 'OWNER_ID': contact_id},
        })
        bound = bitrix.get_all('crm.activity.list', {
            'select': select,
            'filter': {**common_filter, 'BINDINGS': [
                {'OWNER_TYPE_ID': 3, 'OWNER_ID': int(contact_id)},
            ]},
        })
        added = 0
        for activity in [*direct, *bound]:
            activity_id = str(activity.get('ID') or '')
            if not activity_id or activity_id in seen_ids:
                continue
            seen_ids.add(activity_id)
            subject = str(activity.get('SUBJECT') or '')
            if 'Исходящий' not in subject:
                continue
            author_id = str(activity.get('AUTHOR_ID') or '')
            if author_id in EXCLUDED_CALL_AUTHOR_IDS:
                continue
            if author_id not in users:
                users[author_id] = get_record(bitrix, 'user.get', author_id)
            author = users[author_id]
            if author is None:
                raise RuntimeError(f'Звонок {activity_id}: не найден пользователь {author_id}')
            departments = author.get('UF_DEPARTMENT') or []
            if not isinstance(departments, (list, tuple, set)):
                departments = [departments]
            if CONSULTATION_DEPARTMENT_ID not in {str(dep) for dep in departments}:
                continue
            if not activity.get('START_TIME') or not activity.get('END_TIME'):
                continue
            call_start = parse_bitrix_datetime(activity['START_TIME'])
            call_end = parse_bitrix_datetime(activity['END_TIME'])
            if (call_start.year, call_start.month) != (year, month):
                continue
            duration = call_end - call_start
            if duration.total_seconds() < 0:
                raise ValueError(f'Звонок {activity_id}: окончание раньше начала')
            calls.append({
                'date': call_start,
                'contact': name,
                'phone': call_phone(subject),
                'duration': duration,
                'employee': user_name(author, author_id),
            })
            added += 1
        print(f'Контакт {contact_id}: по основному владельцу {len(direct)}, '
              f'по всем привязкам {len(bound)}, в отчёт добавлено {added}')
    return sorted(calls, key=lambda item: item['date'])


def save_workbook(calls, month, year, output_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Звонки ЛК'
    sheet.append([f'Исходящие звонки линии консультаций — {MONTH_NAMES[month - 1]} {year}'])
    sheet.merge_cells('A1:E1')
    sheet.append(['Дата звонка', 'Контакт (кому звонили)',
                  'Номер телефона, на который звонили', 'Хронометраж', 'Кто звонил'])
    for call in calls:
        sheet.append([call['date'].strftime('%d.%m.%Y %H:%M:%S'),
                      call['contact'], call['phone'], call['duration'], call['employee']])
    if not calls:
        sheet.append(['За выбранный период учтённых звонков нет'])
        sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=5)
    total = sum((call['duration'] for call in calls), timedelta())
    sheet.append(['Итого по звонкам', '', '', total])
    thin = Side(style='thin', color='D9E2F3')
    for row in sheet:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.value is not None:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if isinstance(cell.value, timedelta):
                cell.number_format = '[h]:mm:ss'
    for row_no, color, font_color in ((1, '2F5597', 'FFFFFF'),
                                      (2, '2F5597', 'FFFFFF'),
                                      (sheet.max_row, 'E2F0D9', '000000')):
        for cell in sheet[row_no]:
            cell.fill = PatternFill(fill_type='solid', fgColor=color)
            cell.font = Font(bold=True, color=font_color)
    for col, width in {'A': 22, 'B': 32, 'C': 42, 'D': 18, 'E': 26}.items():
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'A3'
    sheet.auto_filter.ref = f'A2:E{max(2, sheet.max_row - 1)}'
    sheet.sheet_view.showGridLines = False
    workbook.save(output_path)
    workbook.close()
    return total


def create_contact_consultation_calls_report(req):
    """Точка входа для маршрута кастомного вебхука; пишет файл в Диск Б24."""
    contact_ids = parse_bp_contact_ids(req['contacts'])
    month, year, start, end = parse_bp_period(req)
    user_id = parse_user_id(req['user_id']) if req.get('user_id') else None

    bitrix = Bitrix(authentication('Bitrix'))
    calls = collect_calls(bitrix, contact_ids, month, year, start, end)
    report_name = (
        f'Звонки_ЛК_контакты_{year}-{month:02d}_'
        f'{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx'
    )
    with tempfile.TemporaryDirectory(prefix='consultation_calls_') as temp_dir:
        report_path = Path(temp_dir) / report_name
        total = save_workbook(calls, month, year, report_path)
        file_content = base64.b64encode(report_path.read_bytes()).decode('ascii')

    upload_result = bitrix.call('disk.folder.uploadfile', {
        'id': BITRIX_REPORT_FOLDER_ID,
        'data': {'NAME': report_name},
        'fileContent': file_content,
    })
    if isinstance(upload_result, dict) and isinstance(upload_result.get('result'), dict):
        upload_result = upload_result['result']
    detail_url = (upload_result.get('DETAIL_URL') or upload_result.get('detailUrl')
                  if isinstance(upload_result, dict) else None)
    if not detail_url:
        raise RuntimeError('Битрикс24 не вернул ссылку на загруженный отчёт')

    if user_id is not None:
        bitrix.call('im.notify.system.add', {
            'USER_ID': user_id,
            'MESSAGE': (f'Отчёт по звонкам ЛК за {MONTH_NAMES[month - 1]} {year} '
                        f'сформирован. Звонков: {len(calls)}. {detail_url}'),
        })
    return {'url': detail_url, 'calls': len(calls),
            'seconds': int(total.total_seconds()), 'contact_ids': contact_ids}


def main():
    contact_ids = parse_contact_ids(input('ID контактов через запятую: '))
    month, year, start, end = parse_period(input('Месяц и год (например, 09 2026): '))
    bitrix = Bitrix(authentication('Bitrix'))
    calls = collect_calls(bitrix, contact_ids, month, year, start, end)
    filename = f'Звонки_ЛК_контакты_{year}-{month:02d}_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
    output_path = Path.cwd() / filename
    total = save_workbook(calls, month, year, output_path)
    seconds = int(total.total_seconds())
    print(f'Звонков: {len(calls)}; всего: {seconds // 3600}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}')
    print(f'Файл: {output_path}')


if __name__ == '__main__':
    main()
